import re, os, glob, subprocess, sys
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, Cm
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from lxml import etree

# ---------------------------------------------------------
# TỰ ĐỘNG CÀI ĐẶT DEPENDENCIES BỊ THIẾU
# ---------------------------------------------------------
def ensure_dependencies():
    try:
        import latex2mathml.converter
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "latex2mathml"])

ensure_dependencies()
import latex2mathml.converter as l2m

# ---------------------------------------------------------
# HỆ THỐNG PIPELINE LÕI (4 TẦNG)
# ---------------------------------------------------------

# TẦNG 1: SANITIZER
def _sanitize_emojis(doc):
    EMOJI_RE = re.compile("[\U00010000-\U0010ffff\u2600-\u27bf\u2300-\u23ff]")
    UGLY_SYMBOLS = {'\uf0b7': '\u2022', '\uf0a7': '\u2022', '\uf076': '-', '\uf0d8': '\u2022'}
    for p in doc.paragraphs:
        for r in p.runs:
            if not r.text: continue
            for bad, good in UGLY_SYMBOLS.items():
                r.text = r.text.replace(bad, good)
            r.text = EMOJI_RE.sub('', r.text)
            r.text = r.text.replace('(\u2022)', '').replace('( )', '').replace('()', '').replace('(:', ':').replace('( :', ':')

# TẦNG 2: SCIENTIFIC & LATEX ENGINE
_xsl_cache = None
def get_mml2omml_xsl():
    global _xsl_cache
    if _xsl_cache is not None: return _xsl_cache
    candidates = []
    for base in (r"C:\Program Files\Microsoft Office", r"C:\Program Files (x86)\Microsoft Office"):
        for ver in ("root\\Office16", "Office16", "root\\Office15", "Office15", "Office14"):
            candidates.append(os.path.join(base, ver, "MML2OMML.XSL"))
    for c in candidates:
        if os.path.isfile(c):
            _xsl_cache = etree.XSLT(etree.parse(c))
            return _xsl_cache
    found = glob.glob(r"C:\Program Files*\Microsoft Office\**\MML2OMML.XSL", recursive=True)
    if found: 
        _xsl_cache = etree.XSLT(etree.parse(found[0]))
        return _xsl_cache
    return None

def latex_to_omml(latex_str, display=False):
    xsl = get_mml2omml_xsl()
    if not xsl: return None
    s = latex_str.strip().replace(r"\dfrac", r"\frac").replace(r"\tfrac", r"\frac")
    s = s.replace(r"^{\circ}", "°").replace(r"^\circ", "°").replace(r"\degree", "°")
    s = s.replace(r"\%", "%")
    try:
        mml = l2m.convert(s, display="block" if display else "inline")
        root = xsl(etree.fromstring(mml.encode("utf-8"))).getroot()
        if etree.QName(root).localname == "oMathPara":
            inner = root.find(f"{{{root.nsmap['m']}}}oMath")
            root = inner if inner is not None else root
        return parse_xml(etree.tostring(root))
    except Exception:
        return None

def _process_latex(doc):
    LATEX_RE = re.compile(r"\$\$(.+?)\$\$|\$(.+?)\$")
    for p in doc.paragraphs:
        text = p.text
        if '$' not in text: continue
        
        matches = list(LATEX_RE.finditer(text))
        if not matches: continue
        
        p.clear()
        last_idx = 0
        for match in matches:
            start, end = match.span()
            if start > last_idx:
                r = p.add_run(text[last_idx:start])
                r.font.name = "Times New Roman"
            
            latex_code = match.group(1) or match.group(2)
            is_display = match.group(1) is not None
            
            # Xử lý Mã nguồn Tin học (Consolas Fallback) nếu chứa keyword code
            is_code = any(kw in latex_code for kw in ['int ', 'float ', 'cout', 'cin', 'boolean'])
            if is_code:
                r = p.add_run(latex_code)
                r.font.name = "Consolas"
                r.font.size = Pt(11)
            else:
                omml_elem = latex_to_omml(latex_code, is_display)
                if omml_elem is not None:
                    p._p.append(omml_elem)
                else:
                    # Fallback cơ bản
                    r = p.add_run(latex_code)
                    r.font.name = "Times New Roman"
                    r.font.italic = True
                
            last_idx = end
        
        if last_idx < len(text):
            r = p.add_run(text[last_idx:])
            r.font.name = "Times New Roman"

# TẦNG 3: EXAM LAYOUT CLASSIFIER
def _classify_and_layout(doc):
    PART_RE = re.compile(r"^(PH\u1ea6N|CH\u01af\u01a0NG|B\u00c0I)\s+([IVXLCDM\d]+)", re.IGNORECASE)
    QUESTION_RE = re.compile(r"^(C\u00e2u|B\u00e0i)\s+([A-Za-z0-9_]+(?:.*?:|.*?\.)\s*)", re.IGNORECASE)
    OPTION_RE = re.compile(r"^([A-D])[\.\)]\s+(.*)$")
    TRUE_FALSE_RE = re.compile(r"^([a-d])[\.\)]\s+(.*)$")
    BULLET_RE = re.compile(r"^([\u2022\-\+o\u25e6])\s+(.*)$")

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text: continue

        # Format chuẩn chung cho toàn bộ doc
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.widow_control = True
        
        for r in p.runs:
            if r.text and r.font.name != "Consolas":
                r.font.name = 'Times New Roman'
                if not r.font.size: r.font.size = Pt(13)

        # 1. PHẦN / CHƯƠNG
        if PART_RE.match(text):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            for r in p.runs: 
                r.font.bold = True
                r.text = r.text.upper()
            continue

        # 2. Câu hỏi (Câu X:)
        q_match = QUESTION_RE.match(text)
        if q_match:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.keep_with_next = True
            label_len = len(q_match.group(0))
            current_len = 0
            for r in p.runs:
                if r.text:
                    if current_len < label_len: r.font.bold = True
                    current_len += len(r.text)
            continue

        # 3. Phương án A, B, C, D
        opt_match = OPTION_RE.match(text)
        if opt_match:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = -Inches(0.25)
            label_len = len(opt_match.group(1)) + 1
            current_len = 0
            for r in p.runs:
                if r.text:
                    if current_len < label_len + 1: r.font.bold = True
                    current_len += len(r.text)
            continue

        # 4. Phương án Đúng/Sai (a, b, c, d)
        tf_match = TRUE_FALSE_RE.match(text)
        if tf_match:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = -Inches(0.25)
            label_len = len(tf_match.group(1)) + 1
            current_len = 0
            for r in p.runs:
                if r.text:
                    if current_len < label_len + 1: r.font.bold = True
                    current_len += len(r.text)
            continue

        # 5. Bullet (- hoặc •)
        if BULLET_RE.match(text) or text.startswith("- "):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = -Inches(0.25)
            continue

        # 6. Văn bản thông thường
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# TẦNG 4: DOCUMENT BUILDER (SETUP PAGE)
def _setup_page(doc, position, show_on_first_page):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.5)
        
        section.different_first_page_header_footer = not show_on_first_page
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if position == "bottom_center" else WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("Trang ")
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        fld_xml = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        r._r.append(fld_xml)


# ---------------------------------------------------------
# INTERFACE CHO DOCX_ENGINE.PY (MASTER CONTROLLER)
# ---------------------------------------------------------

def setup_page_numbering(doc, position="bottom_center", show_on_first_page=False):
    """
    Điểm truy cập duy nhất (Entry point).
    Gom toàn bộ 4 tầng xử lý vào hàm này để đảm bảo thứ tự Pipeline chuẩn tuyệt đối.
    """
    _sanitize_emojis(doc)         # Tầng 1
    _process_latex(doc)           # Tầng 2
    _classify_and_layout(doc)     # Tầng 3
    _setup_page(doc, position, show_on_first_page) # Tầng 4

# Các hàm phụ trợ được mock rỗng để tránh lỗi khi docx_engine.py gọi đến
def format_headings_and_structure(doc): pass
def format_bullet_hierarchy(doc): pass
def convert_markdown_tables(doc): pass
def sanitize_emojis_and_badges(doc): pass
def sanitize_vietnamese_typography(doc): pass
def format_exam_questions(doc): pass
def format_figures_and_captions(doc): pass
def apply_flow_control(doc): pass
