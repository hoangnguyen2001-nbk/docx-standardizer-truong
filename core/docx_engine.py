import io
from docx import Document
from core.engine_utils import TextProcessor, OmmlConverter
from processors import (
    science_math,
    science_chemistry,
    advanced_formatter
)

def process_document(file_bytes, options):
    doc = Document(io.BytesIO(file_bytes))
    
    # UNIFIED STANDARD WORKFLOW (No options needed)
    
    # 1. Page Formatting (1" Left, 0.5" Right/Top/Bottom)
    advanced_formatter.setup_page_numbering(doc)
    
    # 2. General structural formatting (Headings, Standard Paragraphs: 1.25 spacing, 0pt before/after)
    advanced_formatter.format_headings_and_structure(doc)
    
    # 3. List & Bullets (Strict dash '-' with 0.25" hanging indent)
    advanced_formatter.format_bullet_hierarchy(doc)
    
    # 4. Tables (Smart styling without breaking)
    advanced_formatter.convert_markdown_tables(doc)
    
    # 5. Clean Symbols & Typography
    advanced_formatter.sanitize_emojis_and_badges(doc)
    advanced_formatter.sanitize_vietnamese_typography(doc)
    
    # 6. Exam & Figure Captions
    advanced_formatter.format_exam_questions(doc)
    advanced_formatter.format_figures_and_captions(doc)
    
    # 7. Document flow (Widow/Orphan)
    advanced_formatter.apply_flow_control(doc)
    
    # 8. Science & Math Processing (Protect equations, format formulas)
    rules = []
    rules.append('MATH')
    rules.extend(science_math.get_rules())
    rules.extend(science_chemistry.get_rules())
    
    conv = OmmlConverter()
    processor = TextProcessor(rules, conv)
    
    for p in doc.paragraphs:
        processor.process_paragraph(p._p)
        
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    processor.process_paragraph(p._p)
                    
    out_stream = io.BytesIO()
    doc.save(out_stream)
    return out_stream.getvalue()
